from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_dark_modern_presentation():
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Paleta Dark Mode (Cyan y Morado neon sobre oscuro)
    COLOR_BG_OVERLAY = RGBColor(20, 20, 25)      # Oscuro base para superposiciones
    COLOR_ACCENT_1 = RGBColor(0, 210, 255)       # Cyan Neón
    COLOR_ACCENT_2 = RGBColor(138, 43, 226)      # Morado
    COLOR_TEXT_LIGHT = RGBColor(240, 240, 240)   # Texto principal
    COLOR_TEXT_MUTED = RGBColor(170, 170, 180)   # Texto secundario

    bg_image_path = "dark_bg.jpg"

    def set_slide_background(slide):
        # En python-pptx agregar una imagen como fondo real es complejo en algunos layouts,
        # la forma más robusta y compatible es agregar una imagen estirada al tamaño de la slide en el fondo (z-index 0)
        slide.shapes.add_picture(bg_image_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    # =====================================================================
    # Diapositiva 1: El Propósito (El Problema y la Solución)
    # =====================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Título principal sutil
    title_box1 = slide1.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf1 = title_box1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "TRANSFORMACIÓN DE VALOR"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ACCENT_1
    p1.font.name = 'Segoe UI'

    # Bloque 1: Datos Crudos
    shape1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(3), Inches(2.5), Inches(1.5))
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = COLOR_BG_OVERLAY
    shape1.line.color.rgb = COLOR_TEXT_MUTED
    tf_s1 = shape1.text_frame
    tf_s1.word_wrap = True
    p_s1 = tf_s1.paragraphs[0]
    p_s1.text = "DATOS\nCRUDOS"
    p_s1.alignment = PP_ALIGN.CENTER
    p_s1.font.size = Pt(20)
    p_s1.font.bold = True
    p_s1.font.color.rgb = COLOR_TEXT_LIGHT
    p_s1.font.name = 'Segoe UI'

    # Flecha 1
    arrow1 = slide1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.2), Inches(3.5), Inches(0.8), Inches(0.5))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = COLOR_ACCENT_2
    arrow1.line.fill.background()

    # Bloque 2: Motor / Proceso
    shape2 = slide1.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(4.2), Inches(2.75), Inches(2.0), Inches(2.0))
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = COLOR_ACCENT_2
    shape2.line.color.rgb = COLOR_ACCENT_2
    tf_s2 = shape2.text_frame
    tf_s2.word_wrap = True
    p_s2 = tf_s2.paragraphs[0]
    p_s2.text = "MOTOR\nB.I."
    p_s2.alignment = PP_ALIGN.CENTER
    p_s2.font.size = Pt(22)
    p_s2.font.bold = True
    p_s2.font.color.rgb = COLOR_TEXT_LIGHT
    p_s2.font.name = 'Segoe UI'

    # Flecha 2
    arrow2 = slide1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.4), Inches(3.5), Inches(0.8), Inches(0.5))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = COLOR_ACCENT_1
    arrow2.line.fill.background()

    # Bloque 3: Decisiones
    shape3 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(3), Inches(2.5), Inches(1.5))
    shape3.fill.solid()
    shape3.fill.fore_color.rgb = COLOR_BG_OVERLAY
    shape3.line.color.rgb = COLOR_ACCENT_1
    tf_s3 = shape3.text_frame
    tf_s3.word_wrap = True
    p_s3 = tf_s3.paragraphs[0]
    p_s3.text = "DECISIONES\nRENTABLES"
    p_s3.alignment = PP_ALIGN.CENTER
    p_s3.font.size = Pt(20)
    p_s3.font.bold = True
    p_s3.font.color.rgb = COLOR_TEXT_LIGHT
    p_s3.font.name = 'Segoe UI'


    # =====================================================================
    # Diapositiva 2: La Metodología (Los 5 Pasos)
    # =====================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)

    # Título superior
    title_box2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1.2))
    tf_title2 = title_box2.text_frame
    p_title2 = tf_title2.paragraphs[0]
    p_title2.text = "ESTRUCTURA METODOLÓGICA"
    p_title2.font.bold = True
    p_title2.font.size = Pt(36)
    p_title2.font.color.rgb = COLOR_TEXT_LIGHT
    p_title2.font.name = 'Segoe UI'

    # Línea decorativa debajo del título
    line2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.5), Inches(2.0), Inches(0.05))
    line2.fill.solid()
    line2.fill.fore_color.rgb = COLOR_ACCENT_1
    line2.line.fill.background()

    steps = [
        ("01", "DIAGNÓSTICO", "Auditoría de bases de datos y mapeo de procesos."),
        ("02", "ESTRUCTURACIÓN", "Diseño dimensional y pipeline ETL sin código."),
        ("03", "FORMULACIÓN", "Modelado matemático de KPIs operativos y financieros."),
        ("04", "DESARROLLO", "Despliegue del Tablero de Control interactivo."),
        ("05", "EVALUACIÓN", "Análisis de viabilidad económica (VAN, TIR y ROI).")
    ]

    top_start = 2.0
    step_height = 1.0

    for i, (num, title, desc) in enumerate(steps):
        # Número
        num_shape = slide2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(top_start + i * step_height), Inches(0.6), Inches(0.6))
        num_shape.fill.solid()
        num_shape.fill.fore_color.rgb = COLOR_ACCENT_2
        num_shape.line.fill.background()
        tf_num = num_shape.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = num
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.bold = True
        p_num.font.size = Pt(16)
        p_num.font.color.rgb = COLOR_TEXT_LIGHT

        # Título del paso
        title_box = slide2.shapes.add_textbox(Inches(1.3), Inches(top_start + i * step_height - 0.1), Inches(3.0), Inches(0.5))
        tf_step = title_box.text_frame
        p_step = tf_step.paragraphs[0]
        p_step.text = title
        p_step.font.bold = True
        p_step.font.size = Pt(18)
        p_step.font.color.rgb = COLOR_ACCENT_1
        p_step.font.name = 'Segoe UI'

        # Descripción del paso
        desc_box = slide2.shapes.add_textbox(Inches(1.3), Inches(top_start + i * step_height + 0.3), Inches(8.0), Inches(0.5))
        tf_desc = desc_box.text_frame
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = desc
        p_desc.font.size = Pt(16)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.font.name = 'Segoe UI'


    # =====================================================================
    # Diapositiva 3: El Ancla (Las Contribuciones)
    # =====================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)

    # Título
    title_box3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf_title3 = title_box3.text_frame
    p_title3 = tf_title3.paragraphs[0]
    p_title3.text = "CONTRIBUCIONES DEL PROYECTO"
    p_title3.alignment = PP_ALIGN.CENTER
    p_title3.font.bold = True
    p_title3.font.size = Pt(38)
    p_title3.font.color.rgb = COLOR_TEXT_LIGHT
    p_title3.font.name = 'Segoe UI'

    # Línea decorativa
    line3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(1.3), Inches(3.0), Inches(0.05))
    line3.fill.solid()
    line3.fill.fore_color.rgb = COLOR_ACCENT_2
    line3.line.fill.background()

    contributions = [
        "Arquitectura de\nDatos Saneada",
        "Motor de KPIs\nOperativos",
        "Viabilidad Financiera\nDemostrada"
    ]

    card_width = 2.6
    card_height = 4.0
    start_left = 0.8
    spacing = 0.6

    for i, contrib in enumerate(contributions):
        left = start_left + i * (card_width + spacing)

        # Tarjeta semi-transparente simulada (Fondo oscuro)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(2.2), Inches(card_width), Inches(card_height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_OVERLAY
        card.line.color.rgb = COLOR_ACCENT_1

        # Número gigante
        num_box = slide3.shapes.add_textbox(Inches(left), Inches(2.5), Inches(card_width), Inches(1))
        tf_num_card = num_box.text_frame
        p_num_card = tf_num_card.paragraphs[0]
        p_num_card.text = f"0{i+1}"
        p_num_card.alignment = PP_ALIGN.CENTER
        p_num_card.font.bold = True
        p_num_card.font.size = Pt(48)
        p_num_card.font.color.rgb = COLOR_ACCENT_2
        p_num_card.font.name = 'Segoe UI'

        # Texto de la contribución
        text_box = slide3.shapes.add_textbox(Inches(left + 0.1), Inches(4.0), Inches(card_width - 0.2), Inches(1.5))
        tf_text = text_box.text_frame
        tf_text.word_wrap = True
        p_text = tf_text.paragraphs[0]
        p_text.text = contrib
        p_text.alignment = PP_ALIGN.CENTER
        p_text.font.bold = True
        p_text.font.size = Pt(22)
        p_text.font.color.rgb = COLOR_TEXT_LIGHT
        p_text.font.name = 'Segoe UI'

    prs.save('Presentacion_BI.pptx')
    print("Archivo 'Presentacion_BI.pptx' actualizado con tema oscuro, imagen de fondo y diseño profesional.")

if __name__ == '__main__':
    create_dark_modern_presentation()
